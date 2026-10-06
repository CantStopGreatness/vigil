import io

import pytest
from docx import Document

from vigil.export import download_name, executive_summary, render_docx, render_pdf
from vigil.models import Finding, ScanResult, Severity

HOSTILE = '<img src="/etc/passwd"/><font name="x">&amp; \x00\x07 😀 ünïcødé'


def result(findings, passed=(), grade="F", score=0):
    return ScanResult(target="https://example.com/", final_url="https://example.com/",
                      started_at="2026-10-05T12:30:00+00:00", duration_ms=10, score=score, grade=grade,
                      findings=list(findings), passed=list(passed), errors=[])


FINDINGS = [
    Finding("Exposed files", "Environment file exposed", Severity.CRITICAL, "A .env file is public.",
            "Remove `.env` from the web root.", evidence="/.env (HTTP 200) DB_PASSWORD=****\\nKEY=****"),
    Finding("Information disclosure", "Server software versions disclosed", Severity.LOW, "Banner.",
            "Set `server_tokens off;`.", evidence="Server: " + HOSTILE + "x" * 600),
    Finding("Security headers", "Missing Permissions-Policy", Severity.INFO, "Optional.", "Add it."),
]


def docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


@pytest.mark.parametrize("r", [result(FINDINGS), result([], passed=["Cookies", "TLS / certificate"], grade="A",
                                                                   score=100)])
def test_pdf_renders(r):
    pdf = render_pdf(r)
    assert pdf.startswith(b"%PDF") and len(pdf) > 2000


def test_docx_contains_report_content_and_escapes_hostile_text():
    text = docx_text(render_docx(result(FINDINGS)))
    for expected in ("Website Security Report", "Executive summary", "Environment file exposed",
                     "Immediately", "HOW TO FIX", "Methodology and scope"):
        assert expected in text
    assert '<img src="/etc/passwd"/>' in text  # shown literally, not interpreted
    assert "\x00" not in text and "DB_PASSWORD=****\nKEY=****" in text


def test_executive_summary_is_plain_english():
    s = " ".join(executive_summary(result(FINDINGS)))
    assert "scored 0 out of 100 (grade F)" in s
    assert "2 issues: 1 critical and 1 low" in s
    assert "1 informational suggestion." in s
    assert "Fix first: Environment file exposed." in s
    clean = " ".join(executive_summary(result([], passed=["Cookies"], grade="A", score=100)))
    assert "found no security issues" in clean and "Cookies" in clean


def test_download_name_is_safe():
    r = result([])
    r.final_url = 'https://Exa"mple.com/'
    assert download_name(r, "pdf") == "vigil-report-exa-mple.com-2026-10-05.pdf"
